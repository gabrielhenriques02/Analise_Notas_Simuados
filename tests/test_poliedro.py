"""Leitura das planilhas de resultado do Sistema Poliedro."""

from __future__ import annotations

import pytest
from sqlalchemy import select

from app.config import config
from app.db import Sessao
from app.ingest import poliedro
from app.models import Aluno, ClassificacaoPoliedro, Fase

PASTA = config.dir_poliedro
PRIMEIRA = PASTA / "Resultado - 1a Fase - Ciclo 1.xlsx"
FINAL = PASTA / "Resultado Final - Ciclo 1.xlsx"

requer_poliedro = pytest.mark.skipif(
    not PRIMEIRA.exists(), reason="planilhas do Poliedro ausentes (fora do versionamento)"
)


@pytest.fixture(scope="module")
def s():
    sessao = Sessao()
    yield sessao
    sessao.rollback()
    sessao.close()


@requer_poliedro
def test_descobre_ciclo_e_fase_pelo_nome():
    assert poliedro.detectar(PRIMEIRA) == (1, Fase.PRIMEIRA)
    assert poliedro.detectar(FINAL) == (1, Fase.SEGUNDA)


@requer_poliedro
def test_le_a_classificacao_da_1a_fase():
    resultado = poliedro.ler(PRIMEIRA)
    assert len(resultado.linhas) == 35
    murilo = next(l for l in resultado.linhas if "MURILO" in l.nome_normalizado)
    assert murilo.posicao_geral == 8          # o relatório de referência diz "8º"
    assert murilo.posicao_unidade == 1
    assert murilo.nota == 8.61
    assert murilo.aprovado is True


@requer_poliedro
def test_le_a_classificacao_final():
    resultado = poliedro.ler(FINAL)
    assert len(resultado.linhas) == 38
    murilo = next(l for l in resultado.linhas if "MURILO COSER" in l.nome_normalizado)
    assert murilo.posicao_geral == 84         # o relatório de referência diz "84º"
    assert murilo.aprovado is True
    assert sum(1 for l in resultado.linhas if l.aprovado) == 1


@requer_poliedro
def test_cabecalho_e_achado_pelos_rotulos():
    """Ancorar numa linha fixa quebraria: o bloco de título tem altura diferente
    entre os dois arquivos (linha 10 num, linha 9 no outro)."""
    import openpyxl

    for caminho in (PRIMEIRA, FINAL):
        wb = openpyxl.load_workbook(caminho, data_only=True)
        try:
            achado = poliedro._achar_cabecalho(wb[wb.sheetnames[0]])
            assert achado is not None
            _, rotulos = achado
            assert "ALUNO" in rotulos and "RM" in rotulos
        finally:
            wb.close()


@requer_poliedro
def test_importa_casando_pelos_apelidos(s):
    """Regressão: o Poliedro escreve GONÇAVELS e ABREU DE LIMA, e a comparação
    direta de nomes deixava esses dois fora do ranking em silêncio."""
    resumo = poliedro.importar(s, PRIMEIRA)
    assert resumo.gravados == 35
    assert resumo.desconhecidos == []

    guardadas = s.scalars(
        select(ClassificacaoPoliedro).where(
            ClassificacaoPoliedro.ciclo == 1, ClassificacaoPoliedro.fase == Fase.PRIMEIRA
        )
    ).all()
    assert len(guardadas) == 35

    murilo = s.scalar(select(Aluno).where(Aluno.nome_normalizado == "MURILO COSER ROCHA"))
    dele = next(c for c in guardadas if c.aluno_id == murilo.id)
    assert dele.posicao == 8
    s.rollback()


@requer_poliedro
def test_importar_duas_vezes_nao_duplica(s):
    poliedro.importar(s, FINAL)
    antes = s.query(ClassificacaoPoliedro).count()
    poliedro.importar(s, FINAL)
    assert s.query(ClassificacaoPoliedro).count() == antes
    s.rollback()


def test_arquivo_sem_as_colunas_esperadas_avisa(tmp_path):
    import openpyxl

    caminho = tmp_path / "Resultado - 1a Fase - Ciclo 9.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "nada a ver"
    wb.save(caminho)

    resultado = poliedro.ler(caminho)
    assert resultado.linhas == []
    assert any("ALUNO e RM" in a for a in resultado.avisos)
