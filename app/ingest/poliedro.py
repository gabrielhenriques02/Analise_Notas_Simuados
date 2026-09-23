"""Leitura das planilhas de resultado do Sistema Poliedro.

Sao dois arquivos por ciclo:
  Resultado - 1a Fase - Ciclo N.xlsx  -> aba "Classificação Madan"
  Resultado Final - Ciclo N.xlsx      -> aba "Madan Modelo ITA"

O cabecalho e procurado pelos rotulos, nao fixado numa linha: o Poliedro entrega o
arquivo com um bloco de titulo em cima que muda de altura entre os ciclos, e ancorar
numa linha fixa quebraria no primeiro arquivo diferente.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from app.ingest.normalize import normalizar_nome, normalizar_rm
from app.models import Fase

# Rotulos que identificam a linha de cabecalho de cada formato.
ANCORAS = ("ALUNO", "RM")

CICLO_NO_NOME = re.compile(r"ciclo\s*(\d+)", re.I)
PRIMEIRA_FASE_NO_NOME = re.compile(r"1[ªa]?\s*fase", re.I)


@dataclass
class LinhaPoliedro:
    nome_bruto: str
    nome_normalizado: str
    rm: str | None
    posicao_geral: int | None
    posicao_unidade: int | None
    nota: float | None
    aprovado: bool | None = None


@dataclass
class Resultado:
    arquivo: str
    ciclo: int | None
    fase: Fase
    linhas: list[LinhaPoliedro] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    @property
    def base_ranking(self) -> int | None:
        """Maior posicao vista serve de piso para a base de classificados do ciclo."""
        posicoes = [l.posicao_geral for l in self.linhas if l.posicao_geral]
        return max(posicoes) if posicoes else None


def _inteiro(valor) -> int | None:
    try:
        return int(round(float(valor)))
    except (TypeError, ValueError):
        return None


def _decimal(valor) -> float | None:
    try:
        return round(float(valor), 4)
    except (TypeError, ValueError):
        return None


def _sim_nao(valor) -> bool | None:
    if valor is None:
        return None
    texto = normalizar_nome(valor)
    if texto.startswith("SIM"):
        return True
    if texto.startswith("NAO"):
        return False
    return None


def _achar_cabecalho(ws) -> tuple[int, dict[str, int]] | None:
    """Procura a linha que tem ALUNO e RM e devolve o mapa rotulo -> coluna.

    Os dois formatos tem colunas com o mesmo nome em posicoes diferentes
    ("Unid" e sala e unidade; "Geral" e classificacao e nota), entao o mapa guarda
    a primeira ocorrencia e o chamador usa a posicao relativa quando precisa.
    """
    for numero, linha in enumerate(ws.iter_rows(min_row=1, max_row=30), start=1):
        rotulos = {}
        for celula in linha:
            if isinstance(celula.value, str) and celula.value.strip():
                rotulos.setdefault(normalizar_nome(celula.value), celula.column)
        if all(ancora in rotulos for ancora in ANCORAS):
            return numero, rotulos
    return None


def _coluna_de_classificacao(ws, linha_cabecalho: int, rotulos: dict[str, int]) -> tuple[int | None, int | None]:
    """Devolve (coluna geral, coluna unidade) da classificacao.

    "Geral" aparece duas vezes nos arquivos finais — uma na classificacao, outra na
    nota. A da classificacao e a que fica sob o grupo "Classificação", a esquerda da
    coluna do aluno.
    """
    coluna_aluno = rotulos.get("ALUNO")
    geral = unidade = None
    for celula in ws[linha_cabecalho]:
        if not isinstance(celula.value, str):
            continue
        rotulo = normalizar_nome(celula.value)
        if coluna_aluno and celula.column >= coluna_aluno:
            continue                      # depois do aluno ja e bloco de notas
        if rotulo == "GERAL":
            geral = celula.column
        elif rotulo in {"UNID", "UNIDADE"}:
            unidade = celula.column
    return geral, unidade


def _coluna_de_nota(ws, linha_cabecalho: int, rotulos: dict[str, int]) -> int | None:
    """A nota final: "FINAL" na 1ª fase, "GERAL" depois das notas no resultado final."""
    coluna_aluno = rotulos.get("ALUNO", 0)
    if "FINAL" in rotulos and rotulos["FINAL"] > coluna_aluno:
        return rotulos["FINAL"]
    for celula in reversed(list(ws[linha_cabecalho])):
        if isinstance(celula.value, str) and normalizar_nome(celula.value) == "GERAL":
            if celula.column > coluna_aluno:
                return celula.column
    return None


def _coluna_de_aprovacao(ws, linha_cabecalho: int, fase: Fase) -> int | None:
    alvo = "APROVADO 2 FASE" if fase is Fase.SEGUNDA else "APROVADO PARA 2 FASE"
    for celula in ws[linha_cabecalho]:
        if not isinstance(celula.value, str):
            continue
        rotulo = normalizar_nome(celula.value)
        if rotulo.startswith("APROVADO"):
            if fase is Fase.SEGUNDA and "2" in rotulo and "PARA" not in rotulo:
                return celula.column
            if fase is Fase.PRIMEIRA:
                return celula.column
    return None


def detectar(caminho: Path) -> tuple[int | None, Fase]:
    """Ciclo e fase a partir do nome do arquivo."""
    nome = caminho.name
    casamento = CICLO_NO_NOME.search(nome)
    ciclo = int(casamento.group(1)) if casamento else None
    fase = Fase.PRIMEIRA if PRIMEIRA_FASE_NO_NOME.search(nome) else Fase.SEGUNDA
    return ciclo, fase


def ler(caminho: Path | str, *, ciclo: int | None = None, fase: Fase | None = None) -> Resultado:
    caminho = Path(caminho)
    ciclo_detectado, fase_detectada = detectar(caminho)
    ciclo = ciclo if ciclo is not None else ciclo_detectado
    fase = fase or fase_detectada

    resultado = Resultado(arquivo=caminho.name, ciclo=ciclo, fase=fase)
    if ciclo is None:
        resultado.avisos.append(
            "Não consegui descobrir o ciclo pelo nome do arquivo — informe na tela."
        )

    wb = openpyxl.load_workbook(caminho, data_only=True)
    try:
        for aba in wb.sheetnames:
            ws = wb[aba]
            achado = _achar_cabecalho(ws)
            if achado is None:
                continue
            linha_cabecalho, rotulos = achado

            coluna_geral, coluna_unidade = _coluna_de_classificacao(ws, linha_cabecalho, rotulos)
            coluna_nota = _coluna_de_nota(ws, linha_cabecalho, rotulos)
            coluna_aprovado = _coluna_de_aprovacao(ws, linha_cabecalho, fase)
            coluna_aluno, coluna_rm = rotulos["ALUNO"], rotulos["RM"]

            for linha in ws.iter_rows(min_row=linha_cabecalho + 1, values_only=False):
                nome = linha[coluna_aluno - 1].value
                if not isinstance(nome, str) or not nome.strip():
                    continue
                resultado.linhas.append(
                    LinhaPoliedro(
                        nome_bruto=nome.strip(),
                        nome_normalizado=normalizar_nome(nome),
                        rm=normalizar_rm(linha[coluna_rm - 1].value),
                        posicao_geral=_inteiro(linha[coluna_geral - 1].value) if coluna_geral else None,
                        posicao_unidade=_inteiro(linha[coluna_unidade - 1].value) if coluna_unidade else None,
                        nota=_decimal(linha[coluna_nota - 1].value) if coluna_nota else None,
                        aprovado=_sim_nao(linha[coluna_aprovado - 1].value) if coluna_aprovado else None,
                    )
                )
            if resultado.linhas:
                break

        if not resultado.linhas:
            resultado.avisos.append(
                "Não encontrei uma tabela com as colunas ALUNO e RM neste arquivo."
            )
        return resultado
    finally:
        wb.close()


# ── gravacao ────────────────────────────────────────────────────────────────────

@dataclass
class ResumoPoliedro:
    arquivo: str
    ciclo: int | None
    fase: Fase
    gravados: int = 0
    desconhecidos: list[str] = field(default_factory=list)
    base_ranking: int | None = None
    avisos: list[str] = field(default_factory=list)


def importar(session, caminho: Path | str, *, ciclo: int | None = None, fase=None) -> ResumoPoliedro:
    """Grava a classificacao, casando os alunos pelo RM e, na falta dele, pelo nome."""
    from sqlalchemy import select

    from app.models import Aluno, ClassificacaoPoliedro

    leitura = ler(caminho, ciclo=ciclo, fase=fase)
    resumo = ResumoPoliedro(
        arquivo=leitura.arquivo, ciclo=leitura.ciclo, fase=leitura.fase,
        base_ranking=leitura.base_ranking, avisos=list(leitura.avisos),
    )
    if leitura.ciclo is None or not leitura.linhas:
        return resumo

    ativos = list(session.scalars(select(Aluno).where(Aluno.ativo.is_(True))))
    por_rm = {a.rm: a for a in ativos if a.rm}
    por_nome = {a.nome_normalizado: a for a in ativos}

    # O Poliedro escreve os nomes com a mesma bagunca da planilha de correcao
    # (GONÇAVELS, ABREU DE LIMA): sem passar pela tabela de apelidos, dois alunos
    # ficavam de fora do ranking em silencio.
    from app.ingest.normalize import Reconciliador
    from app.models import ApelidoAluno

    apelidos = {
        apelido.nome_normalizado: apelido.aluno.nome_normalizado
        for apelido in session.scalars(select(ApelidoAluno))
    }
    reconciliador = Reconciliador(
        [a.nome_canonico for a in ativos],
        apelidos,
        {a.nome_normalizado: a.nome_canonico for a in ativos},
    )

    for linha in leitura.linhas:
        # O RM e o identificador estavel; o nome varia de grafia entre as fontes.
        aluno = por_rm.get(linha.rm) if linha.rm else None
        if aluno is None:
            canonico = reconciliador.resolver(linha.nome_normalizado)
            aluno = por_nome.get(normalizar_nome(canonico)) if canonico else None
        if aluno is None:
            resumo.desconhecidos.append(linha.nome_bruto)
            continue

        registro = session.scalar(
            select(ClassificacaoPoliedro).where(
                ClassificacaoPoliedro.aluno_id == aluno.id,
                ClassificacaoPoliedro.ciclo == leitura.ciclo,
                ClassificacaoPoliedro.fase == leitura.fase,
            )
        )
        if registro is None:
            registro = ClassificacaoPoliedro(
                aluno_id=aluno.id, ciclo=leitura.ciclo, fase=leitura.fase
            )
            session.add(registro)
        registro.posicao = linha.posicao_geral
        registro.nota = linha.nota
        registro.base_ranking = leitura.base_ranking
        registro.aprovado = linha.aprovado
        resumo.gravados += 1

    session.flush()
    if resumo.desconhecidos:
        resumo.avisos.append(
            f"{len(resumo.desconhecidos)} aluno(s) do arquivo não estão na turma: "
            + ", ".join(sorted(resumo.desconhecidos)[:5])
            + ("…" if len(resumo.desconhecidos) > 5 else "")
        )
    return resumo
