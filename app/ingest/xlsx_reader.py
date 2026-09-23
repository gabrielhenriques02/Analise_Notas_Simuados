"""Leitura das abas da planilha de correcao.

Trata as armadilhas documentadas no CLAUDE.md: linhas fantasma, ruido de float,
vocabulario divergente de nivel e as duas tabelas independentes em NÍVEL SIMULADOS.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from app.ingest.normalize import normalizar_nome, normalizar_rm

ABA_ROSTER = "ALUNOS MADAN 2026"
ABA_SEGUNDA_FASE = "CORREÇÃO 2ª 2025"
ABA_PRIMEIRA_FASE = "CORREÇÃO 1ª FASE 2025"
ABA_PORTUGUES = "CORREÇÃO PORTUGUÊS"
ABA_REDACAO = "CORREÇÃO REDAÇÃO"
ABA_NIVEL = "NÍVEL SIMULADOS"
ABA_FRENTE = "QUESTÃO POR FRENTE"

# A 1ª fase tem 48 questoes: 12 por materia, nesta ordem.
BLOCOS_PRIMEIRA_FASE = [
    ("MATEMÁTICA", 1, 12),
    ("FÍSICA", 13, 24),
    ("QUÍMICA", 25, 36),
    ("INGLÊS", 37, 48),
]

# A planilha usa MÉDIO na 2ª fase e MÉDIA na 1ª para o mesmo nivel.
_NIVEIS = {"FACIL": "FÁCIL", "MEDIO": "MÉDIO", "MEDIA": "MÉDIO", "DIFICIL": "DIFÍCIL"}


def normalizar_nivel(bruto: object) -> str | None:
    return _NIVEIS.get(normalizar_nome(bruto))


def arredondar(valor: object) -> float | None:
    """A planilha guarda ruido de float (5.0999999999999996)."""
    if valor is None or isinstance(valor, str):
        return None
    try:
        return round(float(valor), 4)
    except (TypeError, ValueError):
        return None


@dataclass
class LinhaCorrecao:
    """Um lancamento: um aluno numa prova, com a nota de cada questao."""

    nome_bruto: str
    nome_normalizado: str
    rm: str | None
    ciclo: int
    fase: str
    materia: str
    notas: dict[int, float]  # numero da questao (na materia) -> nota
    nota_final: float | None

    @property
    def tudo_zerado(self) -> bool:
        return bool(self.notas) and all(n == 0 for n in self.notas.values())


@dataclass
class Planilha:
    roster: list[str] = field(default_factory=list)
    lancamentos: list[LinhaCorrecao] = field(default_factory=list)
    niveis: dict[tuple[int, str], dict[int, str]] = field(default_factory=dict)
    frentes: dict[tuple[int, str], dict[int, str]] = field(default_factory=dict)

    def ocorrencias_de_nomes(self) -> dict[str, tuple[set[str], int, set[str]]]:
        """Agrupa as grafias encontradas, para alimentar o reconciliador."""
        saida: dict[str, tuple[set[str], int, set[str]]] = {}
        for linha in self.lancamentos:
            grafias, total, abas = saida.get(linha.nome_normalizado, (set(), 0, set()))
            grafias.add(linha.nome_bruto)
            abas.add(f"{linha.fase} {linha.materia}".strip())
            saida[linha.nome_normalizado] = (grafias, total + 1, abas)
        return saida


def _celulas(aba, min_row: int = 2):
    for linha in aba.iter_rows(min_row=min_row, values_only=True):
        if any(v is not None for v in linha):
            yield linha


def ler(caminho: Path) -> Planilha:
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = Planilha()
        planilha.roster = _ler_roster(wb)
        planilha.lancamentos.extend(_ler_segunda_fase(wb))
        planilha.lancamentos.extend(_ler_primeira_fase(wb))
        planilha.lancamentos.extend(_ler_portugues(wb))
        planilha.lancamentos.extend(_ler_redacao(wb))
        planilha.niveis = _ler_niveis(wb)
        planilha.frentes = _ler_frentes(wb)
        return planilha
    finally:
        wb.close()


def _ler_roster(wb) -> list[str]:
    return [
        str(linha[0]).strip()
        for linha in _celulas(wb[ABA_ROSTER])
        if linha[0] and str(linha[0]).strip()
    ]


def _ler_segunda_fase(wb) -> list[LinhaCorrecao]:
    """NOME | TURMA | CICLO | MATÉRIA | Q1..Q10 (0-10) | NOTA FINAL | CLASSIFICAÇÃO."""
    saida = []
    for linha in _celulas(wb[ABA_SEGUNDA_FASE]):
        nome, _turma, ciclo, materia = linha[0], linha[1], linha[2], linha[3]
        # Linhas fantasma (583-667) tem TURMA e formula viva, mas nao tem nome.
        if not nome or not str(nome).strip() or ciclo is None or not materia:
            continue
        notas = {i + 1: (arredondar(linha[4 + i]) or 0.0) for i in range(10)}
        saida.append(
            LinhaCorrecao(
                nome_bruto=str(nome).strip(),
                nome_normalizado=normalizar_nome(nome),
                rm=None,  # esta aba nao tem RM
                ciclo=int(ciclo),
                fase="2ª FASE",
                materia=str(materia).strip().upper(),
                notas=notas,
                nota_final=arredondar(linha[14]) if len(linha) > 14 else None,
            )
        )
    return saida


def _ler_primeira_fase(wb) -> list[LinhaCorrecao]:
    """Coluna3 | Nome | Coluna1 | RM | CICLO | Q1..Q48 (0/1) | ..."""
    saida = []
    for linha in _celulas(wb[ABA_PRIMEIRA_FASE]):
        nome, rm, ciclo = linha[1], linha[3], linha[4]
        if not nome or not str(nome).strip() or ciclo is None:
            continue
        for materia, primeira, ultima in BLOCOS_PRIMEIRA_FASE:
            notas = {}
            for numero in range(primeira, ultima + 1):
                indice = 5 + (numero - 1)
                valor = arredondar(linha[indice]) if indice < len(linha) else None
                notas[numero - primeira + 1] = valor or 0.0
            saida.append(
                LinhaCorrecao(
                    nome_bruto=str(nome).strip(),
                    nome_normalizado=normalizar_nome(nome),
                    rm=normalizar_rm(rm),
                    ciclo=int(ciclo),
                    fase="1ª FASE",
                    materia=materia,
                    notas=notas,
                    nota_final=None,  # calculado, nunca lido da planilha
                )
            )
    return saida


def _ler_portugues(wb) -> list[LinhaCorrecao]:
    """ALUNO | RM | CICLO | Q1..Q15 (0/1) | NOTA FINAL | CLASSIFICAÇÃO."""
    saida = []
    for linha in _celulas(wb[ABA_PORTUGUES]):
        nome, rm, ciclo = linha[0], linha[1], linha[2]
        if not nome or not str(nome).strip() or ciclo is None:
            continue
        notas = {i + 1: (arredondar(linha[3 + i]) or 0.0) for i in range(15)}
        saida.append(
            LinhaCorrecao(
                nome_bruto=str(nome).strip(),
                nome_normalizado=normalizar_nome(nome),
                rm=normalizar_rm(rm),
                ciclo=int(ciclo),
                fase="2ª FASE",
                materia="PORTUGUÊS",
                notas=notas,
                nota_final=arredondar(linha[18]) if len(linha) > 18 else None,
            )
        )
    return saida


def _ler_redacao(wb) -> list[LinhaCorrecao]:
    """NOME | RM | CICLO | NOTA FINAL | CRITÁRIO 1..CRITÉRIO 5 (0-2)."""
    saida = []
    for linha in _celulas(wb[ABA_REDACAO]):
        nome, rm, ciclo = linha[0], linha[1], linha[2]
        if not nome or not str(nome).strip() or ciclo is None:
            continue
        notas = {}
        for i in range(5):
            indice = 4 + i
            valor = arredondar(linha[indice]) if indice < len(linha) else None
            if valor is not None:
                notas[i + 1] = valor
        saida.append(
            LinhaCorrecao(
                nome_bruto=str(nome).strip(),
                nome_normalizado=normalizar_nome(nome),
                rm=normalizar_rm(rm),
                ciclo=int(ciclo),
                fase="2ª FASE",
                materia="REDAÇÃO",
                notas=notas,
                nota_final=arredondar(linha[3]),
            )
        )
    return saida


def _ler_niveis(wb) -> dict[tuple[int, str], dict[int, str]]:
    """Duas tabelas independentes na mesma aba: A1:L19 (2ª fase) e N2:BJ8 (1ª fase)."""
    aba = wb[ABA_NIVEL]
    grade = [list(linha) for linha in aba.iter_rows(values_only=True)]
    saida: dict[tuple[int, str], dict[int, str]] = {}

    # Bloco A (2ª fase): MATÉRIA | PROVA(=ciclo) | QUESTÃO 1..10, a partir da linha 2.
    for linha in grade[1:19]:
        if not linha or not linha[0] or linha[1] is None:
            continue
        materia = str(linha[0]).strip().upper()
        try:
            ciclo = int(linha[1])
        except (TypeError, ValueError):
            continue
        niveis = {}
        for i in range(10):
            nivel = normalizar_nivel(linha[2 + i]) if 2 + i < len(linha) else None
            if nivel:
                niveis[i + 1] = nivel
        if niveis:
            saida[(ciclo, materia)] = niveis

    # Bloco B (1ª fase): cabecalho na linha 2, CICLO na coluna N (indice 13).
    for linha in grade[2:8]:
        if len(linha) <= 13 or linha[13] is None:
            continue
        bruto = str(linha[13]).strip().upper().replace("CICLO", "").strip()
        try:
            ciclo = int(bruto)
        except ValueError:
            continue
        for materia, primeira, ultima in BLOCOS_PRIMEIRA_FASE:
            niveis = {}
            for numero in range(primeira, ultima + 1):
                indice = 14 + (numero - 1)
                nivel = normalizar_nivel(linha[indice]) if indice < len(linha) else None
                if nivel:
                    niveis[numero - primeira + 1] = nivel
            if niveis:
                saida[(ciclo, f"1ª FASE {materia}")] = niveis
    return saida


def _ler_frentes(wb) -> dict[tuple[int, str], dict[int, str]]:
    """CONCATENAR | CICLO | PROVA | Q1..Q48 — a chave util e (ciclo, prova)."""
    saida: dict[tuple[int, str], dict[int, str]] = {}
    for linha in _celulas(wb[ABA_FRENTE]):
        if linha[1] is None or not linha[2]:
            continue
        try:
            ciclo = int(linha[1])
        except (TypeError, ValueError):
            continue
        prova = str(linha[2]).strip().upper()
        frentes = {}
        for i in range(48):
            indice = 3 + i
            valor = linha[indice] if indice < len(linha) else None
            if valor and str(valor).strip():
                frentes[i + 1] = str(valor).strip().upper()
        if not frentes:
            continue
        if prova == "1ª FASE":
            for materia, primeira, ultima in BLOCOS_PRIMEIRA_FASE:
                bloco = {
                    n - primeira + 1: frentes[n] for n in range(primeira, ultima + 1) if n in frentes
                }
                if bloco:
                    saida[(ciclo, f"1ª FASE {materia}")] = bloco
        else:
            saida[(ciclo, prova)] = frentes
    return saida
